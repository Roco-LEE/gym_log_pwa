# -*- coding: utf-8 -*-
"""앱 기록(JSON) → 헬스일지.xlsx [기록] 시트에 행 추가.

serve.py 가 /api/sync 로 받아서 호출하고, 앱에서 내보낸 JSON 파일로도 쓸 수 있음:
    python sync_excel.py 헬스일지_2026-09-18.json
    python sync_excel.py --test 헬스일지_2026-09-18.json   # 헬스일지_테스트.xlsx 에

- 한 세션의 운동 하나 = [기록] 한 줄. A날짜 C운동 E,F워밍업 G~R 1~6세트 W메모 만 쓰고
  나머지(요일·부위·볼륨·1RM 등)는 시트에 이미 깔린 수식이 계산.
- 그날 전체 메모는 그 세션 첫 줄 W열에 "[오늘] ..." 로 들어감.
- 앱에서 운동에 붙여둔 고정 메모(머신 번호 등)는 [운동목록] F열 뒤에 "[폰] ..." 로 붙임.
  원래 적혀 있던 내용은 건드리지 않고, 다시 동기화하면 [폰] 뒤쪽만 갱신됨.
- 같은 세션을 두 번 넣지 않도록 넣은 id 를 synced.json 에 기록.
- 쓰기 전에 백업/ 폴더에 사본을 남김 (최근 10개).
- openpyxl 로 저장하면 Excel 이 서식을 무시하는 문제가 있어서 생성 스크립트와 같은 후처리를 함.
"""
import openpyxl, datetime, json, os, re, shutil, sys, zipfile, glob

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.environ.get("GYMLOG_XLSX") or os.path.join(HERE, "..", "헬스일지.xlsx")   # 다른 위치면 환경변수로
SYNCED = os.path.join(HERE, "synced.json")
BACKUP_DIR = os.path.join(HERE, "..", "백업")
FIRST_ROW = 3       # 헤더 2줄
MAX_SETS = 6
PHONE_MARK = "[폰]"   # [운동목록] 메모에서 앱이 관리하는 구간의 시작 표시


def use_test_file():
    """테스트 모드: 진짜 헬스일지.xlsx 대신 헬스일지_테스트.xlsx 에 씀 (없으면 진짜 걸 복사해서 만듦)."""
    global XLSX, SYNCED, BACKUP_DIR
    test = os.path.join(HERE, "..", "헬스일지_테스트.xlsx")
    if not os.path.exists(test):
        shutil.copy2(XLSX, test)
    XLSX, SYNCED, BACKUP_DIR = test, os.path.join(HERE, "synced_test.json"), os.path.join(HERE, "..", "백업_테스트")
    return test


def load_synced():
    try:
        with open(SYNCED, encoding="utf-8") as f:
            return set(json.load(f))
    except Exception:
        return set()


def save_synced(ids):
    with open(SYNCED, "w", encoding="utf-8") as f:
        json.dump(sorted(ids), f, ensure_ascii=False, indent=1)


def backup():
    os.makedirs(BACKUP_DIR, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    shutil.copy2(XLSX, os.path.join(BACKUP_DIR, f"헬스일지_{stamp}.xlsx"))
    old = sorted(glob.glob(os.path.join(BACKUP_DIR, "헬스일지_*.xlsx")))
    for p in old[:-10]:
        os.remove(p)


def fix_apply_attrs(path):
    """생성 스크립트의 후처리와 동일. cellXfs 에 apply* 속성을 넣어야 Excel 이 서식을 적용함."""
    zin = zipfile.ZipFile(path, "r")
    items = zin.infolist()
    data = {i.filename: zin.read(i.filename) for i in items}
    zin.close()
    st = data["xl/styles.xml"].decode("utf-8")
    m = re.search(r"(<cellXfs[^>]*>)(.*?)(</cellXfs>)", st, re.S)

    def patch(mo):
        tag = mo.group(0)
        he = tag.find(">")
        head, rest = tag[:he], tag[he + 1:]
        selfclose = head.rstrip().endswith("/")
        if selfclose:
            head = head.rstrip()[:-1].rstrip()
        for attr, name in (("numFmtId", "applyNumberFormat"), ("fontId", "applyFont"),
                           ("fillId", "applyFill"), ("borderId", "applyBorder")):
            v = re.search(attr + r'="(\d+)"', head)
            if v and v.group(1) != "0" and name not in head:
                head += ' %s="1"' % name
        return head + ("/>" if selfclose else ">") + rest

    body = re.sub(r"<xf\b[^>]*/>|<xf\b.*?</xf>", patch, m.group(2), flags=re.S)
    st = st[:m.start(2)] + body + st[m.end(2):]
    data["xl/styles.xml"] = st.encode("utf-8")
    tmp = path + ".tmp"
    zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
    for i in items:
        zout.writestr(i.filename, data[i.filename])
    zout.close()
    shutil.move(tmp, path)


def _merge_ex_memo(cur, phone):
    """[운동목록] F열: 원래 쓰던 글은 두고 '[폰] ...' 부분만 갈아끼움."""
    base = (cur or "").split(PHONE_MARK)[0].strip()
    phone = (phone or "").strip().replace("\n", " ")
    if not phone:
        return base or None
    return f"{base}  {PHONE_MARK} {phone}" if base else f"{PHONE_MARK} {phone}"


def write_ex_memos(wb, ex_memo):
    """앱의 운동 고정 메모를 [운동목록] F열에 반영. 바뀐 개수를 반환."""
    if not ex_memo:
        return 0
    ws = wb["운동목록"]
    changed = 0
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row, 1).value
        if not name or name not in ex_memo:
            continue
        cell = ws.cell(row, 6)
        new = _merge_ex_memo(cell.value, ex_memo[name])
        if new != cell.value:
            cell.value = new
            changed += 1
    # 목록에 없는 운동이면 맨 아래에 새로 추가
    have = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
    for name, memo in ex_memo.items():
        if name in have or not memo.strip():
            continue
        row = ws.max_row + 1
        while ws.cell(row - 1, 1).value in (None, ""):   # 빈 줄이면 위로 당김
            row -= 1
        ws.cell(row, 1, name)
        ws.cell(row, 6, _merge_ex_memo(None, memo))
        changed += 1
    return changed


def append_sessions(sessions, ex_memo=None):
    """sessions: 앱 DB.sessions 형식. 반환: {"written": [id...], "rows": n, "memos": n, "skipped": [id...]}"""
    synced = load_synced()
    todo = [s for s in sessions
            if s.get("id") and s["id"] not in synced and not s.get("fromExcel") and s.get("entries")]
    if not todo and not ex_memo:
        return {"written": [], "rows": 0, "memos": 0, "skipped": [s.get("id") for s in sessions]}

    wb = openpyxl.load_workbook(XLSX)
    memos = write_ex_memos(wb, ex_memo)
    if not todo and not memos:
        return {"written": [], "rows": 0, "memos": 0, "skipped": [s.get("id") for s in sessions]}
    ws = wb["기록"]
    # 첫 빈 줄: A(날짜)와 C(운동)가 모두 비어 있는 첫 행
    row = FIRST_ROW
    while ws.cell(row, 1).value not in (None, "") or ws.cell(row, 3).value not in (None, ""):
        row += 1

    n = 0
    for s in sorted(todo, key=lambda x: x["date"]):
        d = datetime.datetime.strptime(s["date"], "%Y-%m-%d")
        day_memo = (s.get("memo") or "").strip().replace("\n", " ")
        for e in s["entries"]:
            if not e.get("sets") and not e.get("warm"):
                continue
            ws.cell(row, 1, d)
            ws.cell(row, 3, e["ex"])
            if e.get("warm"):
                ws.cell(row, 5, e["warm"]["kg"])
                ws.cell(row, 6, e["warm"]["reps"])
            for i, st in enumerate(e.get("sets", [])[:MAX_SETS]):
                ws.cell(row, 7 + i * 2, st["kg"])
                ws.cell(row, 8 + i * 2, st["reps"])
            memo = (e.get("memo") or "").strip().replace("\n", " ")
            if len(e.get("sets", [])) > MAX_SETS:
                memo = (memo + " " if memo else "") + f"[{len(e['sets'])}세트 중 6세트까지만 기록]"
            if day_memo:   # 그날 메모는 첫 줄에만
                memo = f"[오늘] {day_memo}" + (" · " + memo if memo else "")
                day_memo = ""
            if memo:
                ws.cell(row, 23, memo)
            row += 1
            n += 1

    backup()
    wb.save(XLSX)          # Excel 에서 파일이 열려 있으면 여기서 PermissionError
    fix_apply_attrs(XLSX)
    synced |= {s["id"] for s in todo}
    save_synced(synced)
    return {"written": [s["id"] for s in todo], "rows": n, "memos": memos,
            "skipped": [s["id"] for s in sessions if s["id"] in synced and s not in todo]}


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--test"]
    if "--test" in sys.argv:
        print("테스트 파일:", use_test_file())
    if not args:
        print(__doc__); sys.exit(1)
    with open(args[0], encoding="utf-8") as f:
        data = json.load(f)
    r = append_sessions(data.get("sessions", data if isinstance(data, list) else []), data.get("exMemo"))
    print(f"{r['rows']}줄 추가 (세션 {len(r['written'])}개), 운동 메모 {r['memos']}개")
